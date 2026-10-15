import pandas as pd
from openpyxl import load_workbook

from validacao import sorteio


def _amostra(n=300):
    return pd.DataFrame(
        {
            "full_name": [f"org{i:03d}/repo" for i in range(n)],
            "html_url": [f"https://github.com/org{i:03d}/repo" for i in range(n)],
            "stars": range(n),
        }
    )


def _pares(nome, n=8):
    datas = [f"2026-0{1 + i // 4}-{10 + i:02d}T00:00:00Z" for i in range(n)]
    lista = [
        {"tag_name": f"v1.{i}.0", "published_at": data,
         "html_url": f"https://github.com/{nome}/releases/tag/v1.{i}.0"}
        for i, data in enumerate(datas)
    ]
    return [(lista[i - 1] if i else None, lista[i]) for i in range(n)]


def test_sorteio_de_repositorios_e_reproduzivel():
    primeiro = sorteio.sortear_repositorios(_amostra())
    segundo = sorteio.sortear_repositorios(_amostra().sample(frac=1, random_state=7))
    assert len(primeiro) == 60
    assert primeiro["full_name"].tolist() == segundo["full_name"].tolist()
    assert primeiro["id_repo"].tolist() == list(range(1, 61))
    assert primeiro["full_name"].is_unique


def test_sorteio_com_amostra_menor_que_60():
    assert len(sorteio.sortear_repositorios(_amostra(10))) == 10


def test_releases_elegiveis_precisam_de_anterior():
    elegiveis = sorteio.releases_elegiveis(_pares("o/r", 3))
    assert [r["tag_name"] for r in elegiveis] == ["v1.1.0", "v1.2.0"]
    assert elegiveis[0]["tag_anterior"] == "v1.0.0"


def test_sorteio_de_releases():
    elegiveis = sorteio.releases_elegiveis(_pares("o/r", 8))
    escolhidas = sorteio.sortear_releases(elegiveis)
    assert len(escolhidas) == 5
    assert escolhidas == sorteio.sortear_releases(list(reversed(elegiveis)))
    datas = [r["published_at"] for r in escolhidas]
    assert datas == sorted(datas)


def test_sorteio_de_releases_com_poucas_elegiveis():
    elegiveis = sorteio.releases_elegiveis(_pares("o/r", 4))
    assert len(sorteio.sortear_releases(elegiveis)) == 3
    assert sorteio.sortear_releases([]) == []


def test_links():
    release = {"tag_name": "@scope/pkg@1.2.0"}
    assert sorteio.link_release("o/r", release) == (
        "https://github.com/o/r/releases/tag/@scope/pkg@1.2.0"
    )
    assert sorteio.link_release("o/r", {"tag_name": "x", "html_url": "u"}) == "u"
    assert sorteio.link_commits("o/r", "v1.0.0", "v1.0.1#a") == (
        "https://github.com/o/r/compare/v1.0.0...v1.0.1%23a"
    )


def test_montar_releases_ouro_numera_e_monta_links():
    repositorios = sorteio.sortear_repositorios(_amostra(3))
    releases_ouro = sorteio.montar_releases_ouro(repositorios, _pares)
    assert len(releases_ouro) == 15
    assert releases_ouro["id_release"].tolist() == list(range(1, 16))
    primeira = releases_ouro.iloc[0]
    assert primeira["link_commits"].endswith(f"/compare/{primeira['tag_anterior']}...{primeira['tag_name']}")
    assert set(releases_ouro["id_repo"]) == {1, 2, 3}


def test_planilha_tem_listas_links_e_nenhum_rotulo(tmp_path):
    repositorios = sorteio.sortear_repositorios(_amostra(3))
    releases_ouro = sorteio.montar_releases_ouro(repositorios, _pares)
    caminho = tmp_path / "rotulos_teste.xlsx"
    sorteio.escrever_planilha(caminho, repositorios, releases_ouro)

    livro = load_workbook(caminho)
    assert livro.sheetnames == ["instrucoes", "repositorios", "releases"]
    repos = livro["repositorios"]
    assert [c.value for c in repos[1]] == [
        "id_repo", "repositorio", "link", "tipo_projeto", "entregas_reais", "observacao"
    ]
    assert repos.max_row == 4
    assert repos["C2"].hyperlink.target == repositorios.iloc[0]["html_url"]
    assert repos["D2"].value is None
    listas = {str(v.sqref): v.formula1 for v in repos.data_validations.dataValidation}
    assert listas["D2:D4"] == '"biblioteca/framework,aplicacao/servico,ferramenta CLI,outro"'
    assert listas["E2:E4"] == '"sim,nao,incerto"'

    rels = livro["releases"]
    assert rels.max_row == 16
    assert rels["G2"].hyperlink.target.startswith("https://github.com/")
    assert all(rels.cell(row=linha, column=8).value is None for linha in range(2, 17))
    assert [str(v.sqref) for v in rels.data_validations.dataValidation] == ["H2:H16"]
    # a planilha nao pode mostrar a saida da heuristica
    cabecalhos = [c.value for c in rels[1]]
    assert not any("heuristica" in str(c) for c in cabecalhos)


def test_planilhas_um_arquivo_por_avaliador(tmp_path):
    caminhos = sorteio.planilhas(str(tmp_path), ("ana", "bia"))
    assert [p.split("rotulos_")[1] for p in caminhos] == ["ana.xlsx", "bia.xlsx"]


CONFIG = """
janela: {inicio: "2025-10-02", fim: "2026-10-02"}
alvo_qualificados: 300
busca: {estrelas_min: 1000, faixas: [[1000, null]]}
filtros: {releases_min: 5, runs_min: 50}
"""


def _preparar_main(monkeypatch, tmp_path, argumentos=()):
    import sys

    from validacao.sorteio import releases as modulo_releases

    monkeypatch.chdir(tmp_path)
    (tmp_path / "config.yaml").write_text(CONFIG, encoding="utf-8")
    (tmp_path / "data").mkdir()
    _amostra(70).to_csv(tmp_path / "data" / "amostra.csv", index=False)
    monkeypatch.setattr(sorteio, "ClienteGitHub", lambda **_: object())
    monkeypatch.setattr(modulo_releases, "buscar_releases", lambda cliente, nome, inicio: nome)
    monkeypatch.setattr(modulo_releases, "montar_pares", lambda nome, inicio, fim: _pares(nome))
    monkeypatch.setattr(sys, "argv", ["validacao.sorteio", *argumentos])


def test_main_gera_csvs_e_planilhas(monkeypatch, tmp_path, capsys):
    _preparar_main(monkeypatch, tmp_path)
    sorteio.main()
    repos = pd.read_csv(tmp_path / "validacao" / "amostra_ouro.csv")
    rels = pd.read_csv(tmp_path / "validacao" / "releases_ouro.csv")
    assert len(repos) == 60
    assert len(rels) == 300
    assert (tmp_path / "validacao" / "rotulos_gabriel.xlsx").exists()
    assert (tmp_path / "validacao" / "rotulos_guilherme.xlsx").exists()
    saida = capsys.readouterr().out
    assert "aviso: a amostra tem so 70" in saida


def test_main_nao_sobrescreve_planilha_preenchida(monkeypatch, tmp_path):
    import pytest

    _preparar_main(monkeypatch, tmp_path)
    (tmp_path / "validacao").mkdir()
    preenchida = tmp_path / "validacao" / "rotulos_gabriel.xlsx"
    preenchida.write_bytes(b"rotulos")
    with pytest.raises(SystemExit):
        sorteio.main()
    assert preenchida.read_bytes() == b"rotulos"


def test_main_sobrescreve_quando_pedido(monkeypatch, tmp_path):
    _preparar_main(monkeypatch, tmp_path, ["--sobrescrever"])
    (tmp_path / "validacao").mkdir()
    (tmp_path / "validacao" / "rotulos_gabriel.xlsx").write_bytes(b"velho")
    sorteio.main()
    assert load_workbook(tmp_path / "validacao" / "rotulos_gabriel.xlsx").sheetnames[0] == "instrucoes"
