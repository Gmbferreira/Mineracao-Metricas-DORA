import argparse
import os
import unicodedata

import pandas as pd
from openpyxl import load_workbook
from statsmodels.stats.inter_rater import aggregate_raters, fleiss_kappa

from validacao.sorteio import AVALIADORES, CORRETIVA, ENTREGAS_REAIS, TIPOS_PROJETO

DIR_VALIDACAO = "validacao"

# (dimensao, aba da planilha, coluna do rotulo, itens esperados, categorias)
DIMENSOES = (
    ("tipo_projeto", "repositorios", 4, 60, TIPOS_PROJETO),
    ("entregas_reais", "repositorios", 5, 60, ENTREGAS_REAIS),
    ("corretiva", "releases", 8, 300, CORRETIVA),
)
COLUNA_ID = 1

# Layout padrao: abas "repositorios" e "releases" (o que `validacao.sorteio` gera).
ABAS_PADRAO = ("repositorios", "releases")

# Layout achatado: uma aba com uma linha por release, repetindo os rotulos do
# repositorio (aceito para planilhas reexportadas). Dimensao -> variantes do nome.
COLUNAS_ACHATADO = {
    "id_release": ("id_release",),
    "id_repo": ("id_repo",),
    "corretiva": ("corretiva", "release_corretiva"),
    "tipo_projeto": ("tipo_projeto",),
    "entregas_reais": ("entregas_reais", "entrega_real", "entrega_reais"),
}


def planilhas(diretorio=DIR_VALIDACAO, avaliadores=AVALIADORES):
    return [os.path.join(diretorio, f"rotulos_{nome}.xlsx") for nome in avaliadores]


def _texto(valor):
    """Rotulo sem acentos, minusculo e com espacos simples, para comparar."""
    if valor is None:
        return ""
    decomposto = unicodedata.normalize("NFKD", str(valor).strip())
    sem_acento = "".join(c for c in decomposto if not unicodedata.combining(c))
    return " ".join(sem_acento.lower().split())


def _indice_coluna(cabecalho, variantes):
    normalizado = {_texto(h): i for i, h in enumerate(cabecalho) if h is not None}
    for variante in variantes:
        if _texto(variante) in normalizado:
            return normalizado[_texto(variante)]
    return None


def _registrar(destino, ident, valor, conflitos, dimensao):
    """Guarda o rotulo de um item; acusa se o mesmo id ganhar dois valores."""
    texto = _texto(valor)
    if texto == "":
        return
    if ident in destino:
        if _texto(destino[ident]) != texto:
            conflitos.setdefault(dimensao, set()).add(ident)
    else:
        destino[ident] = valor


def _ler_padrao(livro):
    rotulos = {}
    for nome, aba, col_rotulo, _, _ in DIMENSOES:
        ws = livro[aba]
        valores = {}
        for linha in ws.iter_rows(min_row=2, min_col=COLUNA_ID, max_col=col_rotulo,
                                  values_only=True):
            ident = linha[COLUNA_ID - 1]
            if ident is None:
                continue
            valores[int(ident)] = linha[col_rotulo - 1]
        rotulos[nome] = pd.Series(valores, dtype=object).sort_index()
    return rotulos


def _ler_achatado(livro, caminho=""):
    """Le o layout de uma linha por release (rotulos de repositorio repetidos)."""
    for aba in livro.sheetnames:
        ws = livro[aba]
        linhas = ws.iter_rows(values_only=True)
        try:
            cabecalho = next(linhas)
        except StopIteration:
            continue
        colunas = {nome: _indice_coluna(cabecalho, variantes)
                   for nome, variantes in COLUNAS_ACHATADO.items()}
        if colunas["id_release"] is None or colunas["corretiva"] is None:
            continue
        tipo_projeto, entregas_reais, corretiva = {}, {}, {}
        conflitos = {}
        for linha in linhas:
            if linha[colunas["id_release"]] is None:
                continue
            release = int(linha[colunas["id_release"]])
            _registrar(corretiva, release, linha[colunas["corretiva"]],
                       conflitos, "corretiva")
            if colunas["id_repo"] is None:
                continue
            id_repo = int(linha[colunas["id_repo"]])
            if colunas["tipo_projeto"] is not None:
                _registrar(tipo_projeto, id_repo, linha[colunas["tipo_projeto"]],
                           conflitos, "tipo_projeto")
            if colunas["entregas_reais"] is not None:
                _registrar(entregas_reais, id_repo, linha[colunas["entregas_reais"]],
                           conflitos, "entregas_reais")
        if conflitos:
            detalhe = "; ".join(f"{dim}: {sorted(ids)}"
                                for dim, ids in sorted(conflitos.items()))
            raise SystemExit(
                f"{caminho}: rotulos de repositorio incompativeis entre as releases "
                f"({detalhe})"
            )
        return {
            "tipo_projeto": pd.Series(tipo_projeto, dtype=object).sort_index(),
            "entregas_reais": pd.Series(entregas_reais, dtype=object).sort_index(),
            "corretiva": pd.Series(corretiva, dtype=object).sort_index(),
        }
    raise SystemExit(
        f"{caminho}: layout nao reconhecido (esperado: abas {ABAS_PADRAO} "
        f"ou uma aba com uma linha por release)"
    )


def ler_planilha(caminho):
    """Le uma planilha de rotulagem -> {dimensao: Series indexada por id}.

    Aceita o layout padrao (abas `repositorios` e `releases`) e o layout
    achatado de uma linha por release, comum em planilhas reexportadas.
    """
    livro = load_workbook(caminho, read_only=True)
    try:
        if all(aba in livro.sheetnames for aba in ABAS_PADRAO):
            return _ler_padrao(livro)
        return _ler_achatado(livro, caminho)
    finally:
        livro.close()


def conferir(rotulos, caminho=""):
    """Erros de completude e de valores de uma planilha; [] quando esta pronta."""
    erros = []
    for nome, _, _, n_esperado, categorias in DIMENSOES:
        permitidos = {_texto(c) for c in categorias}
        serie = rotulos.get(nome, pd.Series(dtype=object))
        preenchidos = [v for v in serie if _texto(v) != ""]
        if len(preenchidos) != n_esperado:
            erros.append(f"{caminho} {nome}: {len(preenchidos)}/{n_esperado} preenchidos")
        fora = sorted({str(v).strip() for v in preenchidos
                       if _texto(v) not in permitidos})
        if fora:
            erros.append(f"{caminho} {nome}: valores fora da lista {fora}")
    return erros


def normalizar(serie, categorias):
    canonicos = {_texto(c): c for c in categorias}
    mapeado = {}
    invalidos = []
    for ident, valor in serie.items():
        if _texto(valor) == "":
            invalidos.append(ident)
            continue
        texto = _texto(valor)
        if texto not in canonicos:
            invalidos.append(ident)
            continue
        mapeado[ident] = canonicos[texto]
    if invalidos:
        raise SystemExit(f"rotulos ausentes/invalidos nos ids: {invalidos}")
    return mapeado


def calcular(rotulos_por_avaliador, dimensao):
    """Kappa de Fleiss, concordancia bruta e divergencias de uma dimensao.

    `rotulos_por_avaliador` mapeia nome do avaliador -> {dimensao: Series}.
    Concordancia bruta = fracao de itens com unanimidade dos avaliadores.
    """
    categorias = next(c[4] for c in DIMENSOES if c[0] == dimensao)
    nomes = sorted(rotulos_por_avaliador)
    valores = {nome: normalizar(rotulos_por_avaliador[nome][dimensao], categorias)
               for nome in nomes}
    ids = sorted(next(iter(valores.values())))
    for nome in nomes:
        if sorted(valores[nome]) != ids:
            raise SystemExit(f"ids divergentes entre planilhas na dimensao {dimensao}")

    codigos = {c: i for i, c in enumerate(categorias)}
    matriz = [[codigos[valores[nome][ident]] for nome in nomes] for ident in ids]
    tabela, _ = aggregate_raters(matriz, n_cat=len(categorias))
    kappa = float(fleiss_kappa(tabela))
    unanimidade = [ident for ident, linha in zip(ids, matriz) if len(set(linha)) == 1]
    bruta = len(unanimidade) / len(ids) if ids else float("nan")
    linhas = [
        {"id": ident, **{nome: valores[nome][ident] for nome in nomes}}
        for ident in ids if ident not in set(unanimidade)
    ]
    return {
        "dimensao": dimensao,
        "itens": len(ids),
        "avaliadores": len(nomes),
        "concordancia_bruta": bruta,
        "kappa": kappa,
        "interpretacao": interpretar(kappa),
        "divergencias": pd.DataFrame(linhas),
    }


def interpretar(kappa):
    """Escala de Landis & Koch (1977), como no enunciado."""
    if kappa < 0:
        return "pobre"
    if kappa <= 0.20:
        return "leve"
    if kappa <= 0.40:
        return "razoavel"
    if kappa <= 0.60:
        return "moderada"
    if kappa <= 0.80:
        return "substancial"
    return "quase perfeita"


def pauta_divergencias(rotulos_por_avaliador, diretorio=DIR_VALIDACAO):
    """Divergencias de todas as dimensoes com repositorio, tag e links de apoio.

    Serve de pauta para a discussao de desempate (issue #12): cada linha traz os
    rotulos de cada avaliador e duas colunas vazias (`consenso`, `motivo`) a
    preencher na reuniao.
    """
    amostra = pd.read_csv(os.path.join(diretorio, "amostra_ouro.csv"))
    releases = pd.read_csv(os.path.join(diretorio, "releases_ouro.csv"))
    contexto = {"tipo_projeto": {}, "entregas_reais": {}, "corretiva": {}}
    for linha in amostra.itertuples(index=False):
        info = {"full_name": linha.full_name, "tag": "",
                "link_release": linha.html_url, "link_commits": ""}
        contexto["tipo_projeto"][linha.id_repo] = info
        contexto["entregas_reais"][linha.id_repo] = info
    for linha in releases.itertuples(index=False):
        contexto["corretiva"][linha.id_release] = {
            "full_name": linha.full_name, "tag": linha.tag_name,
            "link_release": linha.link_release, "link_commits": linha.link_commits}
    linhas = []
    for nome, _, _, _, _ in DIMENSOES:
        divergencias = calcular(rotulos_por_avaliador, nome)["divergencias"]
        for item in divergencias.itertuples(index=False):
            info = contexto[nome][item.id]
            linhas.append({
                "dimensao": nome, "id": item.id,
                "full_name": info["full_name"], "tag": info["tag"],
                "link_release": info["link_release"],
                "link_commits": info["link_commits"],
                **{campo: getattr(item, campo) for campo in item._fields[1:]},
                "consenso": "", "motivo": "",
            })
    return pd.DataFrame(linhas)


def main(diretorio=DIR_VALIDACAO, argv=None):
    parser = argparse.ArgumentParser(
        prog="validacao.concordancia",
        description="Calcula o kappa de Fleiss entre os avaliadores da amostra-ouro.",
    )
    parser.add_argument("--diretorio", default=diretorio)
    parser.add_argument("--pauta", default=None,
                        help="grava a lista de divergencias (com links) em CSV")
    args = parser.parse_args(argv)

    faltando = [c for c in planilhas(args.diretorio) if not os.path.exists(c)]
    if faltando:
        raise SystemExit(
            "Planilhas de rotulagem ausentes: " + ", ".join(faltando)
            + ". Termine a rotulagem (issues #10 e #11) antes de calcular a concordancia."
        )
    todos, erros = {}, []
    for caminho in planilhas(args.diretorio):
        rotulos = ler_planilha(caminho)
        pendentes = conferir(rotulos, os.path.basename(caminho))
        if pendentes:
            erros += pendentes
        else:
            nome = os.path.basename(caminho).removeprefix("rotulos_").removesuffix(".xlsx")
            todos[nome] = rotulos
    if erros:
        raise SystemExit("Rotulagem incompleta:\n- " + "\n- ".join(erros)
                         + "\nComplete as planilhas antes de calcular a concordancia.")

    print(f"avaliadores: {', '.join(sorted(todos))}")
    print(f"{'dimensao':<16} {'itens':>5} {'unanim.':>8} {'kappa':>7}  interpretacao")
    for nome, _, _, _, _ in DIMENSOES:
        r = calcular(todos, nome)
        print(f"{nome:<16} {r['itens']:>5} {r['concordancia_bruta']:>7.1%} "
              f"{r['kappa']:>7.3f}  {r['interpretacao']}")
        div = r["divergencias"]
        if len(div):
            print(f"  divergencias ({len(div)}):")
            for linha in div.itertuples(index=False):
                dados = ", ".join(f"{campo}={valor}" for campo, valor in
                                  zip(linha._fields[1:], linha[1:]))
                print(f"    id {linha.id}: {dados}")

    if args.pauta:
        pauta = pauta_divergencias(todos, args.diretorio)
        pauta.to_csv(args.pauta, index=False)
        print(f"\npauta escrita em {args.pauta} ({len(pauta)} divergencias)")


if __name__ == "__main__":
    main()
