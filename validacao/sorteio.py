import argparse
import os
import urllib.parse

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline import config, releases
from pipeline.cliente import ClienteGitHub

# semente fixa: rodar de novo sobre a mesma amostra sorteia os mesmos
# repositorios e as mesmas releases
SEMENTE = 42
N_REPOSITORIOS = 60
N_RELEASES = 5
AVALIADORES = ("gabriel", "guilherme")

DIR_VALIDACAO = "validacao"
ARQ_AMOSTRA = os.path.join("data", "amostra.csv")
ARQ_REPOS_OURO = os.path.join(DIR_VALIDACAO, "amostra_ouro.csv")
ARQ_RELEASES_OURO = os.path.join(DIR_VALIDACAO, "releases_ouro.csv")

TIPOS_PROJETO = ("biblioteca/framework", "aplicacao/servico", "ferramenta CLI", "outro")
ENTREGAS_REAIS = ("sim", "nao", "incerto")
CORRETIVA = ("sim", "nao")

INSTRUCOES = [
    "Rotulagem da amostra-ouro (seção 6 do enunciado)",
    "",
    "Preencha sozinho e não converse com o colega sobre os rótulos até os dois terminarem.",
    "Não olhe a saída da heurística de release corretiva (CFR b no dataset ou o script de avaliação).",
    "Escolha sempre uma das opções da lista de cada coluna. A coluna observacao é livre.",
    "",
    "Aba repositorios",
    "tipo_projeto: biblioteca/framework = código usado por outros projetos como dependência (ex.: react, fastapi);",
    "  aplicacao/servico = programa usado direto pelo usuário ou implantado como serviço (ex.: vscode, supabase);",
    "  ferramenta CLI = programa de linha de comando (ex.: fzf, lazygit); outro = nenhum dos anteriores.",
    "entregas_reais: as releases do repositório são versões que chegam ao usuário (pacote, binário, imagem)?",
    "  sim = cada release é algo que o usuário instala ou usa; nao = marcos internos, snapshots automáticos",
    "  ou releases que não acompanham o que chega ao usuário; incerto = não deu para decidir.",
    "",
    "Aba releases",
    "corretiva: sim = o objetivo principal da release é corrigir um defeito recente (bug, regressão, falha de segurança);",
    "  nao = traz funcionalidade, melhoria, atualização de rotina, ou mistura em que a correção não é o principal.",
    "  Leia as release notes (link_release) e a lista de commits desde a release anterior (link_commits).",
]


def sortear_repositorios(amostra, n=N_REPOSITORIOS, semente=SEMENTE):
    ordenada = amostra.sort_values("full_name").reset_index(drop=True)
    sorteados = ordenada.sample(n=min(n, len(ordenada)), random_state=semente)
    sorteados = sorteados.reset_index(drop=True)
    sorteados.insert(0, "id_repo", range(1, len(sorteados) + 1))
    return sorteados[["id_repo", "full_name", "html_url"]]


def releases_elegiveis(pares):
    """Releases da janela que tem release anterior, para dar link da lista de commits."""
    return [
        {
            "tag_name": release["tag_name"],
            "tag_anterior": anterior["tag_name"],
            "published_at": release["published_at"],
            "html_url": release.get("html_url"),
        }
        for anterior, release in pares
        if anterior is not None
    ]


def sortear_releases(elegiveis, n=N_RELEASES, semente=SEMENTE):
    if not elegiveis:
        return []
    tabela = pd.DataFrame(elegiveis).sort_values(["published_at", "tag_name"])
    tabela = tabela.reset_index(drop=True)
    escolhidas = tabela.sample(n=min(n, len(tabela)), random_state=semente)
    return escolhidas.sort_values("published_at").to_dict("records")


def _ref(tag):
    return urllib.parse.quote(tag, safe="@/")


def link_release(nome, release):
    return release.get("html_url") or f"https://github.com/{nome}/releases/tag/{_ref(release['tag_name'])}"


def link_commits(nome, anterior, tag):
    return f"https://github.com/{nome}/compare/{_ref(anterior)}...{_ref(tag)}"


def montar_releases_ouro(repositorios, buscar_pares):
    """Sorteia as releases de cada repositorio; `buscar_pares(nome)` da os pares."""
    linhas = []
    for repo in repositorios.to_dict("records"):
        nome = repo["full_name"]
        for release in sortear_releases(releases_elegiveis(buscar_pares(nome))):
            linhas.append(
                {
                    "id_release": len(linhas) + 1,
                    "id_repo": repo["id_repo"],
                    "full_name": nome,
                    "tag_name": release["tag_name"],
                    "tag_anterior": release["tag_anterior"],
                    "published_at": release["published_at"],
                    "link_release": link_release(nome, release),
                    "link_commits": link_commits(nome, release["tag_anterior"],
                                                 release["tag_name"]),
                }
            )
    return pd.DataFrame(linhas)


def _cabecalho(aba, colunas, larguras):
    aba.append(colunas)
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for letra, largura in larguras.items():
        aba.column_dimensions[letra].width = largura
    aba.freeze_panes = "A2"


def _link(celula):
    celula.hyperlink = celula.value
    celula.style = "Hyperlink"


def _lista(aba, opcoes, intervalo):
    validacao = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"',
                               allow_blank=True)
    aba.add_data_validation(validacao)
    validacao.add(intervalo)


def escrever_planilha(caminho, repositorios, releases_ouro):
    livro = Workbook()
    instrucoes = livro.active
    instrucoes.title = "instrucoes"
    for texto in INSTRUCOES:
        instrucoes.append([texto])
    instrucoes.column_dimensions["A"].width = 120

    repos = livro.create_sheet("repositorios")
    _cabecalho(repos, ["id_repo", "repositorio", "link", "tipo_projeto",
                       "entregas_reais", "observacao"],
               {"B": 40, "C": 55, "D": 22, "E": 16, "F": 40})
    for repo in repositorios.to_dict("records"):
        repos.append([repo["id_repo"], repo["full_name"], repo["html_url"], None, None, None])
        _link(repos.cell(row=repos.max_row, column=3))
    fim_repos = max(repos.max_row, 2)
    _lista(repos, TIPOS_PROJETO, f"D2:D{fim_repos}")
    _lista(repos, ENTREGAS_REAIS, f"E2:E{fim_repos}")

    rels = livro.create_sheet("releases")
    _cabecalho(rels, ["id_release", "id_repo", "repositorio", "tag", "publicada_em",
                      "link_release", "link_commits", "corretiva", "observacao"],
               {"C": 34, "D": 26, "E": 22, "F": 30, "G": 30, "H": 12, "I": 40})
    for release in releases_ouro.to_dict("records"):
        rels.append([release["id_release"], release["id_repo"], release["full_name"],
                     release["tag_name"], release["published_at"], release["link_release"],
                     release["link_commits"], None, None])
        _link(rels.cell(row=rels.max_row, column=6))
        _link(rels.cell(row=rels.max_row, column=7))
    _lista(rels, CORRETIVA, f"H2:H{max(rels.max_row, 2)}")
    livro.save(caminho)


def planilhas(diretorio=DIR_VALIDACAO, avaliadores=AVALIADORES):
    return [os.path.join(diretorio, f"rotulos_{nome}.xlsx") for nome in avaliadores]


def main():
    parser = argparse.ArgumentParser(
        prog="validacao.sorteio",
        description="Sorteia a amostra-ouro e gera as planilhas de rotulagem.",
    )
    parser.add_argument("--config", default="config.yaml")
    parser.add_argument("--sobrescrever", action="store_true",
                        help="gera de novo mesmo que as planilhas ja existam")
    args = parser.parse_args()

    existentes = [caminho for caminho in planilhas() if os.path.exists(caminho)]
    if existentes and not args.sobrescrever:
        raise SystemExit(
            "As planilhas ja existem e podem ter rotulos: " + ", ".join(existentes)
            + ". Use --sobrescrever so se tiver certeza."
        )

    cfg = config.carregar(args.config)
    amostra = pd.read_csv(ARQ_AMOSTRA)
    if len(amostra) < 300:
        print(f"aviso: a amostra tem so {len(amostra)} repositorios", flush=True)
    repositorios = sortear_repositorios(amostra)
    cliente = ClienteGitHub(cache_dir=cfg.cache_dir, timeout=cfg.timeout)

    def buscar_pares(nome):
        todas = releases.buscar_releases(cliente, nome, cfg.janela_inicio)
        return releases.montar_pares(todas, cfg.janela_inicio, cfg.janela_fim)

    releases_ouro = montar_releases_ouro(repositorios, buscar_pares)
    os.makedirs(DIR_VALIDACAO, exist_ok=True)
    repositorios.to_csv(ARQ_REPOS_OURO, index=False)
    releases_ouro.to_csv(ARQ_RELEASES_OURO, index=False)
    for caminho in planilhas():
        escrever_planilha(caminho, repositorios, releases_ouro)
    print(
        f"{len(repositorios)} repositorios e {len(releases_ouro)} releases sorteados; "
        f"planilhas: {', '.join(planilhas())}",
        flush=True,
    )


if __name__ == "__main__":
    main()
