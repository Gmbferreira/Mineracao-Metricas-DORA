import pandas as pd
import pytest
from openpyxl import Workbook

from validacao import concordancia


def _caso(pares):
    """Dois avaliadores com os mesmos ids; `pares` = [(rotulo_a, rotulo_b), ...]."""
    return {
        "a": {"corretiva": pd.Series([p[0] for p in pares], dtype=object)},
        "b": {"corretiva": pd.Series([p[1] for p in pares], dtype=object)},
    }


def test_interpretacao_seguindo_landis_e_koch():
    assert concordancia.interpretar(-0.3) == "pobre"
    assert concordancia.interpretar(0.10) == "leve"
    assert concordancia.interpretar(0.21) == "razoavel"
    assert concordancia.interpretar(0.50) == "moderada"
    assert concordancia.interpretar(0.79) == "substancial"
    assert concordancia.interpretar(0.81) == "quase perfeita"
    assert concordancia.interpretar(1.0) == "quase perfeita"


def test_acordo_total_da_kappa_um():
    r = concordancia.calcular(_caso([("sim", "sim")] * 10 + [("nao", "nao")] * 10),
                              "corretiva")
    assert r["kappa"] == pytest.approx(1.0)
    assert r["concordancia_bruta"] == 1.0
    assert len(r["divergencias"]) == 0
    assert r["interpretacao"] == "quase perfeita"


def test_kappa_reproduz_o_calculo_manual():
    # 4 itens: 3 unanimes e 1 divergente -> p_o = 0,75, p_e = 0,53125,
    # kappa = (0,75 - 0,53125) / (1 - 0,53125) = 0,4667 (moderada)
    r = concordancia.calcular(_caso([("sim", "sim"), ("sim", "sim"),
                                     ("nao", "nao"), ("sim", "nao")]), "corretiva")
    assert r["kappa"] == pytest.approx(0.4667, abs=1e-4)
    assert r["concordancia_bruta"] == pytest.approx(0.75)
    assert r["interpretacao"] == "moderada"
    div = r["divergencias"]
    assert list(div["id"]) == [3]
    assert div.iloc[0]["a"] == "sim" and div.iloc[0]["b"] == "nao"


def test_quatro_categorias_usando_aggregate_raters():
    categorias = ["outro", "outro", "aplicacao/servico", "biblioteca/framework"]
    dados = {"a": {"tipo_projeto": pd.Series(categorias, dtype=object)},
             "b": {"tipo_projeto": pd.Series(categorias, dtype=object)}}
    r = concordancia.calcular(dados, "tipo_projeto")
    assert r["itens"] == 4
    assert r["avaliadores"] == 2
    assert r["kappa"] == 1.0


def test_categorias_com_maiuscula_sao_aceitas():
    # "ferramenta CLI" tem maiuscula na lista do sorteio; a comparacao e case-insensitive
    tipos = ["ferramenta CLI"] + ["outro"] * 59
    completos = {"tipo_projeto": pd.Series(tipos, dtype=object),
                 "entregas_reais": pd.Series(["sim"] * 60, dtype=object),
                 "corretiva": pd.Series(["nao"] * 300, dtype=object)}
    assert concordancia.conferir(completos) == []
    dados = {"a": {"tipo_projeto": pd.Series(["ferramenta CLI", "outro"], dtype=object)},
             "b": {"tipo_projeto": pd.Series(["ferramenta cli", "outro"], dtype=object)}}
    r = concordancia.calcular(dados, "tipo_projeto")
    assert r["kappa"] == 1.0


def test_ids_divergentes_entre_planilhas():
    dados = {"a": {"corretiva": pd.Series(["sim", "nao"], index=[1, 2], dtype=object)},
             "b": {"corretiva": pd.Series(["sim", "nao"], index=[1, 3], dtype=object)}}
    with pytest.raises(SystemExit, match="ids divergentes"):
        concordancia.calcular(dados, "corretiva")


def test_conferir_acusa_planilha_incompleta():
    rotulos = {"corretiva": pd.Series(["sim"] * 300, dtype=object),
               "tipo_projeto": pd.Series(["outro"] * 59 + [None], dtype=object),
               "entregas_reais": pd.Series(["sim"] * 60, dtype=object)}
    erros = concordancia.conferir(rotulos, "rotulos_x.xlsx")
    assert any("tipo_projeto: 59/60" in e for e in erros)


def test_conferir_acusa_valor_fora_da_lista():
    rotulos = {"corretiva": pd.Series(["talvez"] + ["sim"] * 299, dtype=object),
               "tipo_projeto": pd.Series(["outro"] * 60, dtype=object),
               "entregas_reais": pd.Series(["sim"] * 60, dtype=object)}
    erros = concordancia.conferir(rotulos, "rotulos_x.xlsx")
    assert any("valores fora da lista" in e for e in erros)


def test_conferir_planilha_pronta_nao_reclama():
    rotulos = {"corretiva": pd.Series(["sim"] * 300, dtype=object),
               "tipo_projeto": pd.Series(["outro"] + [None] * 59, dtype=object),
               "entregas_reais": pd.Series(["sim"] + [None] * 59, dtype=object)}
    erros = [e for e in concordancia.conferir(rotulos) if "corretiva" in e]
    assert erros == []


def _planilha_minima(caminho):
    livro = Workbook()
    repos = livro.active
    repos.title = "repositorios"
    repos.append(["id_repo", "repositorio", "link", "tipo_projeto",
                  "entregas_reais", "observacao"])
    repos.append([1, "o/r", "https://example.com", "outro", "sim", None])
    repos.append([2, "o/r2", "https://example.com", None, None, None])
    rels = livro.create_sheet("releases")
    rels.append(["id_release", "id_repo", "repositorio", "tag", "publicada_em",
                 "link_release", "link_commits", "corretiva", "observacao"])
    rels.append([1, 1, "o/r", "v1", "2025-01-01", "u", "u", "nao", None])
    livro.save(caminho)


def test_ler_planilha(tmp_path):
    caminho = tmp_path / "rotulos_teste.xlsx"
    _planilha_minima(caminho)
    rotulos = concordancia.ler_planilha(str(caminho))
    assert list(rotulos["tipo_projeto"].index) == [1, 2]
    assert rotulos["tipo_projeto"][1] == "outro"
    assert rotulos["tipo_projeto"][2] is None
    assert rotulos["corretiva"][1] == "nao"


def test_main_aborta_quando_planilha_falta(tmp_path):
    with pytest.raises(SystemExit, match="ausentes"):
        concordancia.main(diretorio=str(tmp_path), argv=[])


def test_main_aborta_com_rotulagem_incompleta(tmp_path):
    for nome in ("gabriel", "guilherme"):
        _planilha_minima(tmp_path / f"rotulos_{nome}.xlsx")
    with pytest.raises(SystemExit, match="incompleta"):
        concordancia.main(diretorio=str(tmp_path), argv=[])


def _planilha_achatada(caminho):
    """Layout reexportado: uma aba, uma linha por release, rotulos de repo repetidos."""
    livro = Workbook()
    aba = livro.active
    aba.title = "releases_ouro_preenchido"
    aba.append(["id_release", "id_repo", "full_name", "tipo_projeto", "entrega_real",
                "tag_name", "tag_anterior", "published_at", "release_corretiva",
                "link_release", "link_commits"])
    aba.append([1, 1, "o/r", "biblioteca/framework", "não", "v1", None,
                "2025-01-01", "sim", "u", "u"])
    aba.append([2, 1, "o/r", "biblioteca/framework", "não", "v2", None,
                "2025-01-02", "não", "u", "u"])
    aba.append([3, 2, "o/r2", "ferramenta CLI", "sim", "v1", None,
                "2025-01-03", "nao", "u", "u"])
    livro.save(caminho)


def test_ler_planilha_layout_achatado(tmp_path):
    caminho = tmp_path / "rotulos_teste.xlsx"
    _planilha_achatada(caminho)
    rotulos = concordancia.ler_planilha(str(caminho))
    # rotulos de repositorio vem uma vez por id_repo (nao um por release)
    assert list(rotulos["tipo_projeto"].index) == [1, 2]
    assert rotulos["tipo_projeto"][1] == "biblioteca/framework"
    assert rotulos["tipo_projeto"][2] == "ferramenta CLI"
    # 'entrega_real' cai em 'entregas_reais'
    assert rotulos["entregas_reais"][1] == "não"
    # 'release_corretiva' cai em 'corretiva', indexado por id_release
    assert list(rotulos["corretiva"].index) == [1, 2, 3]
    assert rotulos["corretiva"][2] == "não"


def test_ler_planilha_achatada_acusa_conflito_de_repositorio(tmp_path):
    caminho = tmp_path / "rotulos_teste.xlsx"
    livro = Workbook()
    aba = livro.active
    aba.title = "releases_ouro_preenchido"
    aba.append(["id_release", "id_repo", "tipo_projeto", "entrega_real",
                "release_corretiva"])
    aba.append([1, 1, "outro", "sim", "nao"])
    aba.append([2, 1, "biblioteca/framework", "sim", "nao"])
    livro.save(caminho)
    with pytest.raises(SystemExit, match="incompativeis"):
        concordancia.ler_planilha(str(caminho))


def test_ler_planilha_layout_desconhecido_aborta(tmp_path):
    caminho = tmp_path / "rotulos_teste.xlsx"
    livro = Workbook()
    livro.active.append(["qualquer", "coisa"])
    livro.save(caminho)
    with pytest.raises(SystemExit, match="layout nao reconhecido"):
        concordancia.ler_planilha(str(caminho))


def test_acentos_e_maiusculas_viram_categorias_canonicas():
    assert concordancia.normalizar(
        pd.Series(["não", "Sim", "incerto"], dtype=object),
        ["sim", "nao", "incerto"],
    ) == {0: "nao", 1: "sim", 2: "incerto"}

    completos = {"tipo_projeto": pd.Series(["aplicação/serviço"] + ["outro"] * 59,
                                           dtype=object),
                 "entregas_reais": pd.Series(["sim"] * 60, dtype=object),
                 "corretiva": pd.Series(["não"] * 300, dtype=object)}
    assert concordancia.conferir(completos) == []


def test_pauta_traz_contexto_e_links(tmp_path):
    pd.DataFrame({"id_repo": [1, 2], "full_name": ["o/r", "o/r2"],
                  "html_url": ["ur", "ur2"]}).to_csv(
        tmp_path / "amostra_ouro.csv", index=False)
    pd.DataFrame({"id_release": [1, 2], "id_repo": [1, 1],
                  "full_name": ["o/r", "o/r"], "tag_name": ["v1", "v2"],
                  "tag_anterior": ["", ""], "published_at": ["2025-01-01", "2025-01-02"],
                  "link_release": ["lr1", "lr2"],
                  "link_commits": ["lc1", "lc2"]}).to_csv(
        tmp_path / "releases_ouro.csv", index=False)
    dados = {
        "gabriel": {"tipo_projeto": pd.Series(["outro", "biblioteca/framework"],
                                              index=[1, 2], dtype=object),
                    "entregas_reais": pd.Series(["sim", "nao"], index=[1, 2], dtype=object),
                    "corretiva": pd.Series(["sim", "sim"], index=[1, 2], dtype=object)},
        "guilherme": {"tipo_projeto": pd.Series(["outro", "biblioteca/framework"],
                                                index=[1, 2], dtype=object),
                      "entregas_reais": pd.Series(["sim", "nao"], index=[1, 2], dtype=object),
                      "corretiva": pd.Series(["nao", "sim"], index=[1, 2], dtype=object)},
    }
    pauta = concordancia.pauta_divergencias(dados, diretorio=str(tmp_path))
    assert list(pauta["dimensao"]) == ["corretiva"]
    assert list(pauta["id"]) == [1]
    assert pauta.iloc[0]["full_name"] == "o/r"
    assert pauta.iloc[0]["tag"] == "v1"
    assert pauta.iloc[0]["link_commits"] == "lc1"
    assert pauta.iloc[0]["consenso"] == "" and pauta.iloc[0]["motivo"] == ""


def test_tabela_kappa_tem_uma_linha_por_dimensao():
    serie = lambda valores: pd.Series(valores, index=[1, 2, 3, 4], dtype=object)
    dados = {
        "a": {"tipo_projeto": serie(["outro"] * 2 + ["biblioteca/framework"] * 2),
              "entregas_reais": serie(["sim", "sim", "nao", "nao"]),
              "corretiva": serie(["sim", "sim", "nao", "sim"])},
        "b": {"tipo_projeto": serie(["outro"] * 2 + ["biblioteca/framework"] * 2),
              "entregas_reais": serie(["sim", "sim", "nao", "nao"]),
              "corretiva": serie(["sim", "sim", "nao", "nao"])},
    }
    tabela = concordancia.tabela_kappa(dados).set_index("dimensao")
    assert list(tabela.index) == ["tipo_projeto", "entregas_reais", "corretiva"]
    assert tabela.loc["tipo_projeto", "kappa"] == 1.0
    assert tabela.loc["corretiva", "kappa"] == pytest.approx(0.4667, abs=1e-4)
    assert tabela.loc["corretiva", "concordancia_bruta"] == 0.75
    assert tabela.loc["corretiva", "divergencias"] == 1
    assert tabela.loc["corretiva", "interpretacao"] == "moderada"


def test_main_grava_o_kappa(tmp_path, monkeypatch):
    dados = {"gabriel": {}, "guilherme": {}}
    for nome in dados:
        dados[nome] = {
            "tipo_projeto": pd.Series(["outro"] * 60, index=range(1, 61), dtype=object),
            "entregas_reais": pd.Series(["sim"] * 60, index=range(1, 61), dtype=object),
            "corretiva": pd.Series(["nao"] * 300, index=range(1, 301), dtype=object),
        }
        (tmp_path / f"rotulos_{nome}.xlsx").write_bytes(b"")
    monkeypatch.setattr(concordancia, "ler_planilha",
                        lambda caminho: dados["gabriel" if "gabriel" in str(caminho) else "guilherme"])
    saida = tmp_path / "concordancia.csv"
    concordancia.main(diretorio=str(tmp_path), argv=["--saida", str(saida)])
    tabela = pd.read_csv(saida)
    assert list(tabela["dimensao"]) == ["tipo_projeto", "entregas_reais", "corretiva"]
    assert list(tabela["divergencias"]) == [0, 0, 0]
